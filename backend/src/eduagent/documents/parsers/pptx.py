"""Read native slide text/tables; preserve OMML and unhandled-object markers."""

from eduagent.documents.models import DocumentError, MAX_UNITS, result, unit
from eduagent.documents.preflight import pptx_package


def extract(path, kind='pptx'):
    pptx_package(path)
    from pptx import Presentation
    from lxml.etree import tostring
    presentation = Presentation(path)
    if len(presentation.slides) > MAX_UNITS:
        raise DocumentError('too_many_pages')
    units = []
    for number, slide in enumerate(presentation.slides, 1):
        records, warnings = [], ['pptx_reading_order_and_math_unverified']
        title = slide.shapes.title.text if slide.shapes.title else f'第{number}张幻灯片'

        def walk(shapes, ancestors=()):
            for shape in shapes:
                position = dict(shape_id=shape.shape_id, group_path=list(ancestors),
                                left=shape.left, top=shape.top, width=shape.width, height=shape.height)
                if hasattr(shape, 'shapes'):
                    walk(shape.shapes, (*ancestors, shape.shape_id))
                    continue
                if shape.has_text_frame:
                    records.append(dict(kind='paragraphs', text=shape.text, location=position,
                                        paragraphs=[dict(text=p.text, level=p.level) for p in shape.text_frame.paragraphs]))
                if shape.has_table:
                    cells = []
                    for row, values in enumerate(shape.table.rows):
                        for col, cell in enumerate(values.cells):
                            if not cell.is_spanned:
                                cells.append(dict(row=row, column=col, text=cell.text,
                                                  rowspan=cell.span_height, colspan=cell.span_width))
                    records.append(dict(kind='table', cells=cells, location=position,
                                        text='\n'.join(f"{c['row']+1},{c['column']+1}: {c['text']}" for c in cells)))
                for math in shape.element.iter('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'):
                    records.append(dict(kind='formula_omml', xml=tostring(math, encoding='unicode'), location=position,
                                        text='[Office公式原始结构已保留，需核对]'))
                    warnings.append('equation_needs_review')
                # These objects are never executed or silently described as text.
                if int(shape.shape_type) in (13, 6, 3, 16, 24) or (not shape.has_text_frame and not shape.has_table):
                    records.append(dict(kind='unextracted_object', object_type=str(shape.shape_type), location=position, text=''))
                    warnings.append('image_or_complex_object_not_extracted')

        walk(slide.shapes)
        hidden = slide.element.get('show') == '0'
        if hidden:
            warnings.append('hidden_slide_not_for_retrieval')
        if slide.has_notes_slide:
            warnings.append('speaker_notes_not_extracted')
        text = '\n\n'.join(r['text'] for r in records if r.get('text'))
        units.append(unit(number, 'pptx_slide', title, text,
                          dict(slide_number=number, slide_id=slide.slide_id, hidden=hidden),
                          blocks=records, warnings=sorted(set(warnings))))
    return result(units)
