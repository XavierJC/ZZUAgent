"""Self-created source fixtures; no real student/teacher materials."""

from io import BytesIO


def pdf_bytes(texts=('PCA reduces feature dimensions.',), image=False, encrypted=False):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, NumberObject, ArrayObject, TextStringObject, DecodedStreamObject
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=300, height=220)
        # Explicit ToUnicode mapping exercises Chinese extraction without
        # distributing a third-party font. This is not a visual-accuracy fixture.
        letters = list(dict.fromkeys(text))
        codes = {c: i + 1 for i, c in enumerate(letters)}
        pairs = '\n'.join(f'<{codes[c]:04x}> <{c.encode("utf-16-be").hex()}>' for c in letters)
        cmap = DecodedStreamObject()
        cmap.set_data(('begincmap\n1 begincodespacerange\n<0000> <ffff>\nendcodespacerange\n'
                       + str(len(letters)) + ' beginbfchar\n' + pairs + '\nendbfchar\nendcmap').encode())
        cid = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/CIDFontType0'),
            NameObject('/BaseFont'): NameObject('/STSong-Light'), NameObject('/CIDSystemInfo'): DictionaryObject({
                NameObject('/Registry'): TextStringObject('Adobe'), NameObject('/Ordering'): TextStringObject('GB1'), NameObject('/Supplement'): NumberObject(4)})})
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type0'),
            NameObject('/BaseFont'): NameObject('/STSong-Light'), NameObject('/Encoding'): NameObject('/UniGB-UCS2-H'),
            NameObject('/ToUnicode'): writer._add_object(cmap), NameObject('/DescendantFonts'): ArrayObject([writer._add_object(cid)])})
        resources = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        content = 'BT /F1 12 Tf 20 180 Td <' + ''.join(f'{codes[c]:04x}' for c in text) + '> Tj ET'
        if image:
            picture = DecodedStreamObject()
            picture.set_data(bytes([0, 0, 255]) * 100)
            picture.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'),
                NameObject('/Width'): NumberObject(10), NameObject('/Height'): NumberObject(10),
                NameObject('/ColorSpace'): NameObject('/DeviceRGB'), NameObject('/BitsPerComponent'): NumberObject(8)})
            resources[NameObject('/XObject')] = DictionaryObject({NameObject('/Img'): writer._add_object(picture)})
            content += ' q 60 0 0 40 20 40 cm /Img Do Q'
        stream = DecodedStreamObject(); stream.set_data(content.encode())
        page[NameObject('/Resources')] = resources
        page[NameObject('/Contents')] = writer._add_object(stream)
    if encrypted:
        writer.encrypt('synthetic-password')
    output = BytesIO(); writer.write(output)
    return output.getvalue()


def pptx_bytes():
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.oxml.xmlchemy import OxmlElement
    from lxml.etree import fromstring
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = '主成分分析 PCA'
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = '保留主要变化方向'
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(2)).table
    table.cell(0, 0).text = '合并表头'; table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = '方差'; table.cell(1, 1).text = '特征值'
    shape = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1))
    shape.text = '损失公式'
    shape.element.append(fromstring(b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x</m:t></m:r></m:oMath>'))
    other = presentation.slides.add_slide(presentation.slide_layouts[5])
    other.shapes.title.text = '第二页：模型评估'
    other.element.set('show', '0')
    output = BytesIO(); presentation.save(output)
    return output.getvalue()
